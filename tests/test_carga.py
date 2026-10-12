"""Pruebas unitarias de carga, aisladas de MongoDB y de los datos reales."""

import sqlite3

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.carga import Carga
from src.config import TABLAS_SALIDA


@pytest.fixture
def tablas_ejemplo():
    """Genera seis tablas pequeñas con datos y valores nulos de prueba."""
    return {
        "listings": pd.DataFrame({"id": [1, 2], "name": ["Apartamento", None]}),
        "hosts": pd.DataFrame({"host_id": [11, 12], "host_name": ["Ana", "Luis"]}),
        "amenities": pd.DataFrame({"listing_id": [1, 2], "amenity": ["WiFi", "Cocina"]}),
        "reviews": pd.DataFrame({"id": [101], "comments": ["Bien"]}),
        "calendar": pd.DataFrame({"listing_id": [1, 1], "date": ["2026-10-01", "2026-10-02"]}),
        "calendar_mensual": pd.DataFrame({"listing_id": [1], "mes": [10]}),
    }


def test_carga_sqlite_crea_seis_tablas_y_conserva_nulos(tmp_path, tablas_ejemplo):
    db = tmp_path / "output" / "prueba.db"
    carga = Carga(tablas_ejemplo, sqlite_path=db, xlsx_path=tmp_path / "prueba.xlsx")
    carga.cargar_sqlite()

    with sqlite3.connect(db) as conexion:
        nombres = {
            fila[0]
            for fila in conexion.execute("SELECT name FROM sqlite_master WHERE type='table'")
        }
        assert nombres == set(TABLAS_SALIDA)
        for nombre, df in tablas_ejemplo.items():
            assert conexion.execute(f'SELECT COUNT(*) FROM "{nombre}"').fetchone()[0] == len(df)
        valor_nulo = conexion.execute('SELECT name FROM "listings" WHERE id = 2').fetchone()[0]
        assert valor_nulo is None


def test_xlsx_crea_cinco_hojas_sin_calendar(tmp_path, tablas_ejemplo):
    xlsx = tmp_path / "output" / "prueba.xlsx"
    carga = Carga(tablas_ejemplo, sqlite_path=tmp_path / "prueba.db", xlsx_path=xlsx)
    carga.exportar_xlsx()

    libro = load_workbook(xlsx, read_only=True)
    try:
        assert set(libro.sheetnames) == set(TABLAS_SALIDA) - {"calendar"}
        for nombre in libro.sheetnames:
            assert libro[nombre].max_row == len(tablas_ejemplo[nombre]) + 1
    finally:
        libro.close()


def test_rechaza_falta_de_tabla(tmp_path, tablas_ejemplo):
    del tablas_ejemplo["calendar_mensual"]
    carga = Carga(tablas_ejemplo, sqlite_path=tmp_path / "prueba.db", xlsx_path=tmp_path / "prueba.xlsx")
    with pytest.raises(ValueError, match="faltantes"):
        carga.cargar_todo()


def test_cargar_todo_genera_ambos_archivos(tmp_path, tablas_ejemplo):
    db = tmp_path / "output" / "prueba.db"
    xlsx = tmp_path / "output" / "prueba.xlsx"
    Carga(tablas_ejemplo, sqlite_path=db, xlsx_path=xlsx).cargar_todo()
    assert db.is_file()
    assert xlsx.is_file()
